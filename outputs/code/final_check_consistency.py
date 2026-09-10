# -*- coding: utf-8 -*-
"""最终校核: 结果文件 <-> 文档表格 <-> 模型输出 三者一致性"""
import openpyxl, numpy as np, re, sys, json
sys.path.insert(0, 'work')
doc = 'outputs/药材烘干问题A题_建模与求解说明文档.md'
txt = open(doc).read()


def get_table(title, nxt):
    i = txt.index(title); j = txt.index(nxt, i)
    rows = []
    for line in txt[i:j].split("\n"):
        if not line.startswith("|") or re.match(r"^\|[\s\-\|]+\|$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if not re.match(r"^[\d.]+$", cells[0]):
            continue
        rows.append(cells)
    return rows


wb1 = openpyxl.load_workbook('outputs/result1.xlsx')
checks = [("**表 1", "**表 2", "温度", [0, 0.5, 1.0, 1.5, 2.0], 1.0),
          ("**表 2", "**结果解读", "水分浓度", [0, 0.5, 1.0, 1.5, 2.0], 1.0)]
for title, nxt, sheet, rr_list, dt in checks:
    ws = wb1[sheet]; rows = get_table(title, nxt); err = 0.0; n = 0
    for row in rows:
        tt = float(row[0]); vals = [float(x) for x in row[1:]]
        ir = int(round(tt / dt)) + 2
        for j, rr in enumerate(rr_list):
            ic = int(round(rr / 0.1)) + 2
            err = max(err, abs(ws.cell(ir, ic).value - vals[j])); n += 1
    print("%s vs result1.xlsx[%s]: %d 个格子, 最大偏差 %.2e" % (title, sheet, n, err))

wb2 = openpyxl.load_workbook('outputs/result2.xlsx')
rows = get_table("**表 3", "**表 4")
err = 0.0
for row in rows:
    tt = float(row[0]) * 3600; vals = [float(x) for x in row[1:]]
    ir = int(round(tt / 1.0)) + 2
    for j, rr in enumerate([0, 0.5, 1.0, 1.5, 2.0]):
        ic = int(round(rr / 0.1)) + 2
        err = max(err, abs(wb2["温度"].cell(ir, ic).value - vals[j]))
print("**表 3 vs result2.xlsx[温度]: 最大偏差 %.2e" % err)
rows = get_table("**表 4", "**结果解读")
err = 0.0
for row in rows:
    tt = float(row[0]) * 3600; vals = [float(x) for x in row[1:]]
    ir = int(round(tt / 1.0)) + 2
    for j, rr in enumerate([0, 0.5, 1.0, 1.5, 2.0]):
        ic = int(round(rr / 0.1)) + 2
        err = max(err, abs(wb2["水分浓度"].cell(ir, ic).value - vals[j]))
print("**表 4 vs result2.xlsx[水分浓度]: 最大偏差 %.2e" % err)

# 表5/表6 与 result3/4 在整 6 h 时刻对比
wb3 = openpyxl.load_workbook('outputs/result3.xlsx'); ws3 = wb3.worksheets[0]
rows = get_table("**表 5", "**结果解读")
err = 0.0; n = 0
for row in rows:
    th = float(row[0])
    if abs(th * 3600 - round(th * 3600 / 60) * 60) > 1e-9:
        continue                       # 跳过"烘干结束"插值行
    ir = int(round(th * 3600 / 60.0)) + 2
    for j, rr in enumerate([0, 0.5, 1.0, 1.5, 2.0]):
        ic = int(round(rr / 0.1)) + 2
        err = max(err, abs(ws3.cell(ir, ic).value - float(row[1 + j]))); n += 1
print("**表 5 vs result3.xlsx: %d 个格子, 最大偏差 %.2e" % (n, err))

wb4 = openpyxl.load_workbook('outputs/result4.xlsx'); ws4 = wb4.worksheets[0]
rows = get_table("**表 6", "注：表中")
err = 0.0; n = 0
for row in rows:
    th = float(row[0])
    if abs(th * 3600 - round(th * 3600 / 60) * 60) > 1e-9:
        continue
    ir = int(round(th * 3600 / 60.0)) + 2
    for k, rr in enumerate([0, 0.5, 1.0]):
        ic = int(round(rr / 0.1)) + 2
        err = max(err, abs(ws4.cell(ir, ic).value - float(row[1 + k]))); n += 1
    err = max(err, abs(ws4.cell(ir, ws4.max_column).value - float(row[4]))); n += 1
print("**表 6 vs result4.xlsx: %d 个格子, 最大偏差 %.2e" % (n, err))

fc = json.load(open('work/out/final_check.json'))
print("\n最终高分辨率校核 (N=800, 30 s 输出):")
print("  问题3: %.4f h  (生产 N=400 为 57.1470 h, 偏差 %.4f h = %.3f%%)"
      % (fc['p3']['t_dry_h'], fc['p3']['t_dry_h'] - 57.14704342797384,
         100 * (fc['p3']['t_dry_h'] - 57.14704342797384) / 57.14704342797384))
print("  问题4: %.4f h  (生产 N=400 为 50.8173 h, 偏差 %.4f h = %.3f%%)"
      % (fc['p4']['t_dry_h'], fc['p4']['t_dry_h'] - 50.817311479812854,
         100 * (fc['p4']['t_dry_h'] - 50.817311479812854) / 50.817311479812854))
