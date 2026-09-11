# -*- coding: utf-8 -*-
"""按采纳后的口径重生成 result4.xlsx 与 result4_local.xlsx：
   固定物理距离列 0—2.0 cm（材料外留空）+ "药材表面"列 + "表面位置"工作表（记录 R(t)）"""
import sys, os, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import openpyxl
import herb_v2 as v2

RD = os.environ.get("DELIV") or os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "out", "result_files"))
def collect(solver, res, radii):
    n = len(res["t"]); nr = len(radii) + 1
    C = np.full((n, nr), np.nan)
    for i, tt in enumerate(res["t"]):
        _, Co = solver.profiles_at_radii(tt, res["T"][i], res["C"][i], radii)
        C[i, :-1] = Co; C[i, -1] = res["C"][i, -1]
    return C

def emit(tag, solver, rad, t_dry):
    n = int(np.floor(t_dry / 60.0)) + 1
    te = np.concatenate([np.arange(0, n * 60.0, 60.0), [t_dry]])
    rb = solver.run(t_dry, v2.schedule_prod, t_eval=te)
    R_OUT = np.round(np.arange(0, 2.0001, 0.1), 4) / 100.0
    C = collect(solver, rb, R_OUT)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("水分浓度")
    ws.append(["时间\\到药材中心的距离"] + [round(float(x) * 100.0, 4) for x in R_OUT] + ["药材表面"])
    for i, tt in enumerate(rb["t"]):
        ws.append([float(tt)] + [None if not np.isfinite(v) else round(float(v), 4) for v in C[i]])
    ws2 = wb.create_sheet("表面位置")
    ws2.append(["时间\\实际半径", "实际半径/cm"])
    for tt in rb["t"]:
        ws2.append([float(tt), round(float(rad.R_of(tt)) * 100, 4)])
    wb.save(os.path.join(RD, tag))
    Cmax = rb["C"].max(axis=1); i = np.where(Cmax < 0.15)[0][0]
    print("%s: %d 行 × %d 列(含表面列) + 表面位置表 | 烘干 %.4f h | 严格达标首秒 %d s, Cmax=%.10f"
          % (tag, len(rb["t"]), len(R_OUT) + 2, t_dry / 3600, rb["t"][i], Cmax[i]))
    return float(rb["t"][i]), float(Cmax[i])


if __name__ == "__main__":
    t_oven, T_inf, C_inf = v2.hm.load_oven()
    env = v2.hm.Environment(t_oven, T_inf, C_inf)
    t_rad, R_rad = v2.hm.load_radius()
    rad = v2.hm.RadiusLaw(t_rad, R_rad)
    res = {}
    s0 = v2.LagSolver(v2.PropsV2("p4"), rad, env, N=400, mapping="geometric", cap_mode="mass")
    res["基线"] = emit("result4.xlsx", s0, rad, 50.811872367695265 * 3600)
    s1 = v2.LagSolver(v2.PropsV2("p4"), rad, env, N=400, mapping="local", cap_mode="mass")
    res["非均匀"] = emit("result4_local.xlsx", s1, rad, 47.62075179004256 * 3600)
    json.dump({k: {"t_stop_s": v[0], "Cmax": v[1]} for k, v in res.items()},
              open(os.path.join(os.path.dirname(RD), "v2_stop.json"), "w"), ensure_ascii=False, indent=1)
